import io
import re
import time
from collections import Counter
from typing import List, Dict, Tuple

import numpy as np
import pandas as pd
import streamlit as st

import pdfplumber
from docx import Document

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from sentence_transformers import SentenceTransformer


# ============================================================
# CvScout
# Data Science Resume Screening System
# ============================================================

APP_NAME = "CvScout"
APP_SUBTITLE = "Data Science Resume Screening System"

MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"

# ------------------------------------------------------------
# Page configuration
# ------------------------------------------------------------

st.set_page_config(
    page_title=APP_NAME,
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded"
)


# ------------------------------------------------------------
# Styling
# ------------------------------------------------------------

st.markdown(
    """
    <style>

    .main {
        background-color: #050816;
    }

    .block-container {
        padding-top: 2rem;
        padding-bottom: 3rem;
    }

    h1, h2, h3 {
        color: white;
    }

    .metric-card {
        background: #0d1220;
        border: 1px solid #20283b;
        border-radius: 12px;
        padding: 18px;
    }

    .score-good {
        color: #21c97a;
        font-weight: 700;
    }

    .score-review {
        color: #f2c94c;
        font-weight: 700;
    }

    .score-bad {
        color: #ff5c5c;
        font-weight: 700;
    }

    .small-text {
        color: #9aa4b2;
        font-size: 0.9rem;
    }

    </style>
    """,
    unsafe_allow_html=True
)


# ============================================================
# Skill dictionary
# ============================================================

SKILLS = [
    # Programming
    "Python",
    "R",
    "Java",
    "C++",
    "C#",
    "JavaScript",
    "TypeScript",
    "PHP",
    "Scala",

    # Data Science
    "Pandas",
    "NumPy",
    "Scikit-learn",
    "Scikit Learn",
    "TensorFlow",
    "PyTorch",
    "Keras",
    "Machine Learning",
    "Deep Learning",
    "NLP",
    "Natural Language Processing",
    "Computer Vision",
    "Data Science",
    "Data Analysis",
    "Statistics",
    "Statistical Analysis",
    "Predictive Analytics",
    "Regression",
    "Classification",
    "Clustering",
    "Time Series",

    # BI
    "Power BI",
    "Tableau",
    "Looker",
    "Looker Studio",
    "Excel",
    "VBA",
    "DAX",

    # SQL / Databases
    "SQL",
    "MySQL",
    "PostgreSQL",
    "Oracle",
    "SQL Server",
    "MongoDB",
    "Snowflake",
    "Databricks",

    # Big Data
    "Spark",
    "PySpark",
    "Hadoop",
    "Hive",
    "Airflow",

    # Cloud
    "AWS",
    "Azure",
    "GCP",
    "Google Cloud",

    # DevOps
    "Docker",
    "Kubernetes",
    "Terraform",
    "Jenkins",
    "Git",
    "GitHub",
    "GitLab",
    "Linux",

    # Web
    "HTML",
    "CSS",
    "React",
    "Angular",
    "Node.js",
    "Django",
    "Flask",
    "FastAPI",
    "REST API",

    # Project / Management
    "Agile",
    "Scrum",
    "Project Management",
    "Stakeholder Management",
    "Business Analysis",
    "Business Intelligence",

    # Recruitment
    "Recruitment",
    "Talent Acquisition",
    "Candidate Management",
    "Sourcing",
]


# Canonical skills and common variations. Word boundaries avoid false
# matches such as detecting the skill "R" inside every English word.
SKILL_ALIASES = {
    "Python": ["python"],
    "R": ["r programming", "r language", "r"],
    "Java": ["java"],
    "C++": ["c++", "cpp"],
    "C#": ["c#", "c sharp"],
    "JavaScript": ["javascript", "java script"],
    "TypeScript": ["typescript", "type script"],
    "PHP": ["php"],
    "Scala": ["scala"],
    "Pandas": ["pandas"],
    "NumPy": ["numpy"],
    "Scikit-learn": ["scikit-learn", "scikit learn", "sklearn"],
    "TensorFlow": ["tensorflow"],
    "PyTorch": ["pytorch", "py torch"],
    "Keras": ["keras"],
    "Machine Learning": ["machine learning", "machine-learning", "ml"],
    "Deep Learning": ["deep learning", "deep-learning"],
    "NLP": ["natural language processing", "nlp"],
    "Computer Vision": ["computer vision", "opencv"],
    "Data Science": ["data science", "data scientist"],
    "Data Analysis": ["data analysis", "data analytics", "data analyst", "analytics"],
    "Statistics": ["statistics", "statistical analysis", "statistical modelling", "statistical modeling"],
    "Predictive Analytics": ["predictive analytics", "predictive modelling", "predictive modeling"],
    "Regression": ["regression"],
    "Classification": ["classification"],
    "Clustering": ["clustering"],
    "Time Series": ["time series", "time-series forecasting"],
    "Power BI": ["power bi", "powerbi"],
    "Tableau": ["tableau"],
    "Looker": ["looker"],
    "Looker Studio": ["looker studio", "google data studio"],
    "Excel": ["excel", "microsoft excel", "ms excel"],
    "VBA": ["vba"],
    "DAX": ["dax"],
    "SQL": ["sql", "structured query language"],
    "MySQL": ["mysql", "my sql"],
    "PostgreSQL": ["postgresql", "postgres"],
    "Oracle": ["oracle database", "oracle"],
    "SQL Server": ["sql server", "microsoft sql server", "ms sql server"],
    "MongoDB": ["mongodb", "mongo db"],
    "Snowflake": ["snowflake"],
    "Databricks": ["databricks"],
    "Spark": ["apache spark", "spark"],
    "PySpark": ["pyspark", "py spark"],
    "Hadoop": ["hadoop"],
    "Hive": ["apache hive", "hive"],
    "Airflow": ["apache airflow", "airflow"],
    "AWS": ["amazon web services", "aws"],
    "Azure": ["microsoft azure", "azure"],
    "GCP": ["google cloud platform", "gcp", "google cloud"],
    "Docker": ["docker"],
    "Kubernetes": ["kubernetes", "k8s"],
    "Terraform": ["terraform"],
    "Jenkins": ["jenkins"],
    "Git": ["git"],
    "GitHub": ["github"],
    "GitLab": ["gitlab", "git lab"],
    "Linux": ["linux"],
    "HTML": ["html"],
    "CSS": ["css"],
    "React": ["react", "react.js", "reactjs"],
    "Angular": ["angular", "angular.js", "angularjs"],
    "Node.js": ["node.js", "nodejs", "node js"],
    "Django": ["django"],
    "Flask": ["flask"],
    "FastAPI": ["fastapi", "fast api"],
    "REST API": ["rest api", "restful api", "rest apis"],
    "Agile": ["agile"],
    "Scrum": ["scrum"],
    "Project Management": ["project management", "project manager"],
    "Stakeholder Management": ["stakeholder management", "stakeholder engagement"],
    "Business Analysis": ["business analysis", "business analyst"],
    "Business Intelligence": ["business intelligence", "bi reporting"],
    "Recruitment": ["recruitment", "recruiting"],
    "Talent Acquisition": ["talent acquisition"],
    "Candidate Management": ["candidate management"],
    "Sourcing": ["sourcing", "candidate sourcing"],
}


# ============================================================
# Stopwords
# ============================================================

STOPWORDS = {
    "the", "and", "for", "with", "that", "this", "from", "have",
    "has", "will", "are", "you", "your", "our", "their", "they",
    "them", "who", "what", "when", "where", "which", "into",
    "about", "over", "under", "using", "used", "use", "work",
    "working", "role", "roles", "job", "jobs", "candidate",
    "candidates", "experience", "years", "year", "required",
    "requirements", "requirement", "responsibilities",
    "responsibility", "skills", "skill", "ability", "strong",
    "good", "knowledge", "team", "teams", "join", "company",
    "position", "including", "such", "must", "should", "would",
    "can", "able", "also", "other", "more", "than", "through",
    "within", "across", "based", "provide", "support", "ensure",
    "develop", "development", "manage", "management"
}


# ============================================================
# Model
# ============================================================

@st.cache_resource
def load_embedding_model():
    return SentenceTransformer(MODEL_NAME)


# ============================================================
# Text extraction
# ============================================================

def extract_text_from_pdf(uploaded_file) -> str:

    text_parts = []

    try:
        uploaded_file.seek(0)

        with pdfplumber.open(uploaded_file) as pdf:

            for page in pdf.pages:

                page_text = page.extract_text()

                if page_text:
                    text_parts.append(page_text)

    except Exception as e:
        return f"ERROR_READING_PDF: {e}"

    return "\n".join(text_parts)


def extract_text_from_docx(uploaded_file) -> str:

    try:
        uploaded_file.seek(0)

        document = Document(uploaded_file)

        paragraphs = []

        for paragraph in document.paragraphs:

            if paragraph.text.strip():
                paragraphs.append(paragraph.text)

        # Also read tables
        for table in document.tables:

            for row in table.rows:

                for cell in row.cells:

                    if cell.text.strip():
                        paragraphs.append(cell.text)

        return "\n".join(paragraphs)

    except Exception as e:
        return f"ERROR_READING_DOCX: {e}"


def extract_text(uploaded_file) -> str:

    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)

    if filename.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)

    return ""


# ============================================================
# Text normalization
# ============================================================

def normalize_text(text: str) -> str:

    if not text:
        return ""

    text = text.lower()

    # Normalize common variants
    replacements = {
        "scikit learn": "scikit-learn",
        "machine-learning": "machine learning",
        "data-analytics": "data analytics",
        "powerbi": "power bi",
        "sqlserver": "sql server",
        "nodejs": "node.js",
        "google cloud platform": "google cloud",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    text = re.sub(r"\s+", " ", text)

    return text.strip()


# ============================================================
# Contact extraction
# ============================================================

def extract_email(text: str) -> str:

    pattern = r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"

    match = re.search(pattern, text)

    return match.group(0) if match else ""


def extract_phone(text: str) -> str:

    patterns = [
        r"(?:\+44\s?7\d{3}|\+44\s?\d{4}|07\d{3})\s?\d{3}\s?\d{3}",
        r"\+?\d[\d\s().-]{8,}\d"
    ]

    for pattern in patterns:

        matches = re.findall(pattern, text)

        if matches:

            cleaned = re.sub(r"\s+", " ", matches[0]).strip()

            return cleaned

    return ""


def extract_linkedin(text: str) -> str:

    pattern = r"(https?://)?(www\.)?linkedin\.com/in/[A-Za-z0-9\-_%]+"

    match = re.search(pattern, text, re.IGNORECASE)

    if match:

        return match.group(0)

    return ""


def extract_github(text: str) -> str:

    pattern = r"(https?://)?(www\.)?github\.com/[A-Za-z0-9\-_%]+"

    match = re.search(pattern, text, re.IGNORECASE)

    if match:

        return match.group(0)

    return ""


# ============================================================
# Name extraction
# ============================================================

def extract_candidate_name(text: str, filename: str = "") -> str:

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    ignored = {
        "curriculum vitae",
        "resume",
        "cv",
        "profile",
        "professional profile",
        "personal profile",
        "contact",
        "contact details"
    }

    # First few lines are usually the best place
    for line in lines[:20]:

        clean = re.sub(r"[^A-Za-z .'-]", "", line).strip()

        words = clean.split()

        if 2 <= len(words) <= 5:

            lower = clean.lower()

            if lower in ignored:
                continue

            if any(
                word.lower() in ignored
                for word in words
            ):
                continue

            # Avoid obvious headings
            if clean.upper() == clean and len(clean) > 25:
                continue

            return clean

    # Fallback to filename
    base = re.sub(
        r"\.(pdf|docx)$",
        "",
        filename,
        flags=re.IGNORECASE
    )

    base = re.sub(
        r"[\(\[].*?[\)\]]",
        "",
        base
    )

    base = re.sub(
        r"[-_]+",
        " ",
        base
    )

    base = re.sub(
        r"\b(totaljobs|cv|resume)\b",
        "",
        base,
        flags=re.IGNORECASE
    )

    base = re.sub(r"\s+", " ", base).strip()

    return base.title()


def extract_first_name(name: str) -> str:

    if not name:
        return ""

    return name.split()[0]


# ============================================================
# Experience extraction
# ============================================================

def extract_experience_years(text: str) -> float:

    values = []

    patterns = [
        r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\s*(?:of)?\s*experience",
        r"experience\s*[:\-]?\s*(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)",
        r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)"
    ]

    for pattern in patterns:

        matches = re.findall(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        for match in matches:

            try:
                value = float(match)

                if 0 <= value <= 50:
                    values.append(value)

            except:
                pass

    if not values:
        return 0.0

    return max(values)


def extract_required_experience(text: str) -> float:

    patterns = [
        r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\s*(?:of)?\s*(?:relevant\s*)?experience",
        r"minimum\s*(?:of\s*)?(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)",
        r"at least\s*(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)"
    ]

    values = []

    for pattern in patterns:

        matches = re.findall(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        for match in matches:

            try:
                values.append(float(match))
            except:
                pass

    if not values:
        return 0.0

    return max(values)


def experience_score(
    candidate_years: float,
    required_years: float
) -> float:
    """Estimate experience fit. Missing information receives a neutral score."""
    if required_years <= 0:
        return 60.0

    # A zero here often means the parser could not detect the duration,
    # not that the candidate has no experience.
    if candidate_years <= 0:
        return 55.0

    if candidate_years >= required_years:
        extra = min((candidate_years - required_years) * 3, 15)
        return min(100.0, 85.0 + extra)

    ratio = candidate_years / required_years
    return max(25.0, ratio * 85.0)


# ============================================================
# Location extraction
# ============================================================

UK_CITIES = [
    "London",
    "Birmingham",
    "Manchester",
    "Leeds",
    "Liverpool",
    "Bristol",
    "Sheffield",
    "Nottingham",
    "Leicester",
    "Coventry",
    "Reading",
    "Milton Keynes",
    "Cambridge",
    "Oxford",
    "Southampton",
    "Portsmouth",
    "Brighton",
    "Edinburgh",
    "Glasgow",
    "Cardiff",
    "Newcastle",
    "Sunderland",
    "Belfast",
    "Derby",
    "Bradford",
    "York",
    "Peterborough",
    "Norwich",
    "Exeter",
    "Bath",
    "Slough",
    "Watford",
]


def extract_location(text: str) -> str:

    for city in UK_CITIES:

        if re.search(
            rf"\b{re.escape(city)}\b",
            text,
            flags=re.IGNORECASE
        ):
            return city

    return ""


# ============================================================
# Job title / company
# ============================================================

def extract_job_title(text: str) -> str:

    title_patterns = [
        r"(?:current\s+)?job\s+title\s*[:\-]\s*(.+)",
        r"(?:current\s+)?position\s*[:\-]\s*(.+)",
        r"title\s*[:\-]\s*(.+)"
    ]

    for pattern in title_patterns:

        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        if match:

            value = match.group(1).strip()

            if 1 <= len(value) <= 100:
                return value

    # Scan lines for common professional titles
    title_keywords = [
        "manager",
        "engineer",
        "analyst",
        "developer",
        "consultant",
        "director",
        "administrator",
        "technician",
        "specialist",
        "coordinator",
        "recruiter",
        "scientist",
        "architect"
    ]

    for line in text.splitlines()[:80]:

        line_clean = line.strip()

        if not line_clean:
            continue

        lower = line_clean.lower()

        if any(k in lower for k in title_keywords):

            if 2 <= len(line_clean.split()) <= 10:

                return line_clean

    return ""


def extract_current_company(text: str) -> str:

    patterns = [
        r"current\s+company\s*[:\-]\s*(.+)",
        r"company\s*[:\-]\s*(.+)",
        r"employer\s*[:\-]\s*(.+)"
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        if match:

            value = match.group(1).strip()

            if 1 <= len(value) <= 100:
                return value

    return ""


# ============================================================
# Education
# ============================================================

def extract_education(text: str) -> str:

    education_terms = [
        "phd",
        "doctorate",
        "master of science",
        "msc",
        "master of arts",
        "ma",
        "master of business administration",
        "mba",
        "bachelor of science",
        "bsc",
        "bachelor of arts",
        "ba",
        "bachelor of engineering",
        "beng",
        "be",
        "degree",
        "diploma"
    ]

    found = []

    for line in text.splitlines():

        lower = line.lower()

        if any(term in lower for term in education_terms):

            line_clean = line.strip()

            if line_clean and line_clean not in found:

                found.append(line_clean)

    return " | ".join(found[:4])


def education_score(jd_text: str, cv_text: str) -> float:
    """Compare education levels without punishing missing extraction as a mismatch."""
    jd = normalize_text(jd_text)
    cv = normalize_text(cv_text)

    def level(text: str) -> int:
        if any(term in text for term in ["phd", "ph.d", "doctorate", "doctoral"]):
            return 3
        if any(term in text for term in [
            "master of science", "master of arts", "master's", "masters", "msc", "mba", "master degree"
        ]):
            return 2
        if any(term in text for term in [
            "bachelor", "bsc", "b.sc", "beng", "b.eng", "undergraduate degree", "bachelors"
        ]):
            return 1
        return 0

    # Only compare degree levels when the JD clearly names one.
    required_level = level(jd)
    candidate_level = level(cv)
    if required_level == 0:
        return 60.0
    if candidate_level == 0:
        # Missing education text is unknown, not proof of no qualification.
        return 55.0
    if candidate_level >= required_level:
        return 100.0
    return round(max(20.0, candidate_level / required_level * 80.0), 1)


# ============================================================
# Skill matching
# ============================================================

def _alias_found(text: str, alias: str) -> bool:
    """Match an alias as a phrase/token, not as an arbitrary substring."""
    alias = alias.strip().lower()
    if not alias:
        return False
    # Special case the single-letter language R: only match an isolated R.
    if alias == "r":
        return bool(re.search(r"(?<!\w)r(?!\w)", text, flags=re.IGNORECASE))
    pattern = r"(?<![a-z0-9])" + re.escape(alias) + r"(?![a-z0-9])"
    return bool(re.search(pattern, text, flags=re.IGNORECASE))


def find_skills(text: str) -> List[str]:
    """Return deduplicated canonical skills using safe word-boundary matching."""
    if not text:
        return []
    normalized = normalize_text(text)
    found = []
    # Longer aliases are evaluated first inside each skill's alias list.
    for canonical, aliases in SKILL_ALIASES.items():
        if any(_alias_found(normalized, alias) for alias in aliases):
            found.append(canonical)
    return found


def calculate_skill_score(
    jd_text: str,
    cv_text: str
) -> Tuple[float, List[str], List[str]]:
    """Weighted coverage of JD-mentioned skills; essential skills count more."""
    jd_skills = find_skills(jd_text)
    cv_skill_set = set(find_skills(cv_text))

    if not jd_skills:
        # No skills recognized in the JD: don't pretend the candidate failed.
        return 55.0, [], []

    important_words = re.compile(
        r"\b(required|essential|must[- ]have|mandatory|minimum|required skills|"
        r"need to have|key requirement|must possess)\b",
        flags=re.IGNORECASE
    )
    preferred_words = re.compile(
        r"\b(preferred|desirable|bonus|nice to have|advantageous)\b",
        flags=re.IGNORECASE
    )

    essential_skills = set()
    preferred_skills = set()
    for line in jd_text.splitlines():
        line_skills = set(find_skills(line))
        if not line_skills:
            continue
        if important_words.search(line):
            essential_skills.update(line_skills)
        elif preferred_words.search(line):
            preferred_skills.update(line_skills)

    matched = sorted(set(jd_skills) & cv_skill_set)
    missing = sorted(set(jd_skills) - cv_skill_set)

    weighted_total = 0.0
    weighted_matched = 0.0
    for skill in set(jd_skills):
        weight = 2.0 if skill in essential_skills else (0.6 if skill in preferred_skills else 1.0)
        weighted_total += weight
        if skill in cv_skill_set:
            weighted_matched += weight

    score = 100.0 * weighted_matched / weighted_total if weighted_total else 55.0
    return round(score, 1), matched, missing


# ============================================================
# Keyword extraction
# ============================================================

def extract_keywords(
    text: str,
    max_words: int = 40
) -> List[str]:

    text = normalize_text(text)

    words = re.findall(
        r"\b[a-zA-Z][a-zA-Z0-9+#.\-]{2,}\b",
        text
    )

    words = [
        word.lower()
        for word in words
        if word.lower() not in STOPWORDS
    ]

    counts = Counter(words)

    return [
        word
        for word, _ in counts.most_common(max_words)
    ]


def keyword_coverage_score(
    jd_text: str,
    cv_text: str
) -> Tuple[float, List[str]]:
    """Measure JD keyword coverage across the entire CV, not just its top words."""
    jd_keywords = extract_keywords(jd_text, max_words=30)
    if not jd_keywords:
        return 55.0, []

    normalized_cv = normalize_text(cv_text)
    cv_tokens = set(re.findall(r"[a-zA-Z][a-zA-Z0-9+#.\-]{1,}", normalized_cv))
    matched = [word for word in jd_keywords if word in cv_tokens]
    coverage = len(matched) / max(len(jd_keywords), 1)

    # Keep this as a coverage measure; don't artificially inflate it.
    return round(coverage * 100, 1), matched


# ============================================================
# TF-IDF
# ============================================================

def calculate_tfidf_similarity(
    jd_text: str,
    cv_text: str
) -> float:
    """Calibrated lexical similarity. It is a ranking signal, not a probability."""
    jd = normalize_text(jd_text)
    cv = normalize_text(cv_text)
    if not jd or not cv:
        return 0.0

    try:
        vectorizer = TfidfVectorizer(
            stop_words="english",
            ngram_range=(1, 2),
            min_df=1,
            sublinear_tf=True,
            max_features=25000
        )
        matrix = vectorizer.fit_transform([jd, cv])
        raw = float(cosine_similarity(matrix[0:1], matrix[1:2])[0][0])
        # Map typical low document cosine values onto a more usable score range.
        score = np.clip((raw - 0.02) / 0.28 * 100, 0, 100)
        return round(float(score), 1)
    except ValueError:
        return 0.0


def split_into_chunks(text: str, chunk_words: int = 110, overlap_words: int = 25,
                      max_chunks: int = 35) -> List[str]:
    """Split long documents into overlapping chunks to avoid model truncation."""
    words = re.findall(r"\S+", text or "")
    if not words:
        return []
    if len(words) <= chunk_words:
        return [" ".join(words)]

    step = max(1, chunk_words - overlap_words)
    starts = list(range(0, len(words), step))
    if len(starts) > max_chunks:
        # Sample across the full CV/JD rather than keeping only the beginning.
        indexes = np.linspace(0, len(starts) - 1, max_chunks).round().astype(int)
        starts = [starts[i] for i in sorted(set(indexes))]

    chunks = []
    for start in starts:
        chunk = " ".join(words[start:start + chunk_words])
        if chunk:
            chunks.append(chunk)
        if start + chunk_words >= len(words):
            break
    return chunks


@st.cache_data(show_spinner=False, max_entries=16)
def encode_text_chunks(text: str, model_name: str, chunk_words: int,
                       overlap_words: int, max_chunks: int):
    """Cache chunk embeddings so the same JD is not encoded for every resume."""
    chunks = split_into_chunks(
        text, chunk_words=chunk_words,
        overlap_words=overlap_words, max_chunks=max_chunks
    )
    if not chunks:
        return [], np.empty((0, 384), dtype=np.float32)
    model = load_embedding_model()
    embeddings = model.encode(
        chunks, convert_to_numpy=True, normalize_embeddings=True,
        show_progress_bar=False, batch_size=32
    )
    return chunks, embeddings


def calculate_semantic_similarity(
    jd_text: str,
    cv_text: str,
    model
) -> float:
    """Compare JD and CV chunks so key information beyond the first page counts."""
    if not jd_text.strip() or not cv_text.strip():
        return 0.0

    try:
        jd_chunks, jd_embeddings = encode_text_chunks(
            jd_text, MODEL_NAME, 110, 20, 20
        )
        cv_chunks, cv_embeddings = encode_text_chunks(
            cv_text, MODEL_NAME, 110, 25, 35
        )
        if not jd_chunks or not cv_chunks:
            return 0.0

        similarities = np.matmul(jd_embeddings, cv_embeddings.T)
        # For each JD chunk, keep the strongest related CV chunk. This makes a
        # relevant career section count even when the CV contains unrelated text.
        best_per_jd_chunk = similarities.max(axis=1)
        best_scores = np.sort(best_per_jd_chunk)[-min(3, len(best_per_jd_chunk)):]
        raw = float(0.65 * best_scores.max() + 0.35 * best_scores.mean())

        # Convert cosine similarity to a convenient screening scale; not accuracy.
        score = np.clip((raw - 0.20) / 0.45 * 100, 0, 100)
        return round(float(score), 1)
    except Exception:
        return 0.0


# ============================================================
# ATS score
# ============================================================

def calculate_ats_score(text: str) -> float:

    if not text:
        return 0.0

    score = 0

    # Length
    if len(text) > 1200:
        score += 20

    if len(text) > 2500:
        score += 10

    # Contact information
    if extract_email(text):
        score += 20

    if extract_phone(text):
        score += 10

    if extract_linkedin(text):
        score += 10

    # Career / technical information
    if extract_experience_years(text) > 0:
        score += 10

    if find_skills(text):
        score += 15

    if extract_education(text):
        score += 5

    return min(
        100,
        float(score)
    )


# ============================================================
# Final scoring engine
# ============================================================

def calculate_final_score(
    semantic_score: float,
    tfidf_score: float,
    skill_score: float,
    keyword_score: float,
    experience_score_value: float,
    education_score_value: float
) -> float:
    """Weighted relevance score. The score is not a probability or accuracy."""
    final_score = (
        semantic_score * 0.30 +
        tfidf_score * 0.10 +
        skill_score * 0.25 +
        keyword_score * 0.20 +
        experience_score_value * 0.10 +
        education_score_value * 0.05
    )
    return round(float(np.clip(final_score, 0, 100)), 1)




# ============================================================
# Candidate processor
# ============================================================

def process_candidate(
    uploaded_file,
    jd_text: str,
    model
) -> Dict:

    filename = uploaded_file.name

    cv_text = extract_text(uploaded_file)

    if not cv_text.strip() or cv_text.startswith("ERROR_READING_"):
        error_message = cv_text if cv_text.startswith("ERROR_READING_") else "Could not extract resume text. The PDF may be scanned/image-only."
        return {
            "File Name": filename,
            "Full Name": "",
            "First Name": "",
            "Email": "",
            "Phone": "",
            "City": "",
            "Country": "UK",
            "LinkedIn Link": "",
            "Current Job Title": "",
            "Experience (Years)": 0,
            "Current Company": "",
            "Skills": "",
            "Education": "",
            "Certificates": "",
            "Location": "",
            "ATS Score": 0,
            "Github/Website/Portfolio": "",
            "Matched Skills": "",
            "Missing Skills": "",
            "TF-IDF Similarity": 0,
            "Semantic Similarity": 0,
            "Skill Match": 0,
            "Keyword Match": 0,
            "Experience Match": 0,
            "Education Match": 0,
            "Relevancy (%)": 0,
            "Match Summary": error_message
        }

    # --------------------------------------------------------
    # Extraction
    # --------------------------------------------------------

    name = extract_candidate_name(
        cv_text,
        filename
    )

    first_name = extract_first_name(name)

    email = extract_email(cv_text)

    phone = extract_phone(cv_text)

    linkedin = extract_linkedin(cv_text)

    github = extract_github(cv_text)

    experience = extract_experience_years(cv_text)

    current_company = extract_current_company(cv_text)

    job_title = extract_job_title(cv_text)

    location = extract_location(cv_text)

    education = extract_education(cv_text)

    skills = find_skills(cv_text)

    ats_score = calculate_ats_score(cv_text)

    # --------------------------------------------------------
    # Data Science scoring
    # --------------------------------------------------------

    tfidf_score = calculate_tfidf_similarity(
        jd_text,
        cv_text
    )

    semantic_score = calculate_semantic_similarity(
        jd_text,
        cv_text,
        model
    )

    skill_score, matched_skills, missing_skills = (
        calculate_skill_score(
            jd_text,
            cv_text
        )
    )

    keyword_score, matched_keywords = (
        keyword_coverage_score(
            jd_text,
            cv_text
        )
    )

    required_exp = extract_required_experience(
        jd_text
    )

    exp_score = experience_score(
        experience,
        required_exp
    )

    edu_score = education_score(
        jd_text,
        cv_text
    )

    # --------------------------------------------------------
    # Final score
    # --------------------------------------------------------

    relevancy = calculate_final_score(
        semantic_score,
        tfidf_score,
        skill_score,
        keyword_score,
        exp_score,
        edu_score
    )

    # --------------------------------------------------------
    # Match explanation
    # --------------------------------------------------------

    if matched_skills:

        summary = (
            f"Matched {len(matched_skills)} required skill(s)"
        )

    else:

        summary = "Low direct skill overlap"

    if required_exp > 0:

        summary += (
            f" | Required: {required_exp:g} yrs"
            f" | Candidate: {experience:g} yrs"
        )

    return {
        "File Name": filename,
        "Full Name": name,
        "First Name": first_name,
        "Email": email,
        "Phone": phone,
        "City": location,
        "Country": "UK",
        "LinkedIn Link": linkedin,
        "Current Job Title": job_title,
        "Experience (Years)": experience,
        "Current Company": current_company,
        "Skills": ", ".join(skills),
        "Education": education,
        "Certificates": "",
        "Location": location,
        "ATS Score": ats_score,
        "Github/Website/Portfolio": github,
        "Matched Skills": ", ".join(matched_skills),
        "Missing Skills": ", ".join(missing_skills),
        "TF-IDF Similarity": tfidf_score,
        "Semantic Similarity": semantic_score,
        "Skill Match": round(skill_score, 1),
        "Keyword Match": round(keyword_score, 1),
        "Experience Match": round(exp_score, 1),
        "Education Match": round(edu_score, 1),
        "Relevancy (%)": relevancy,
        "Match Summary": summary
    }


# ============================================================
# Export
# ============================================================

EXPORT_COLUMNS = [
    "Relevancy (%)",
    "File Name",
    "Full Name",
    "First Name",
    "Email",
    "Phone",
    "City",
    "Country",
    "LinkedIn Link",
    "Current Job Title",
    "Experience (Years)",
    "Current Company",
    "Skills",
    "Education",
    "Certificates",
    "Location",
    "ATS Score",
    "Github/Website/Portfolio",
    "Matched Skills",
    "Missing Skills",
    "TF-IDF Similarity",
    "Semantic Similarity",
    "Skill Match",
    "Keyword Match",
    "Experience Match",
    "Education Match",
    "Match Summary"
]


def create_excel(df: pd.DataFrame) -> bytes:

    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            index=False,
            sheet_name="CvScout Results"
        )

    output.seek(0)

    return output.getvalue()


# ============================================================
# Sidebar
# ============================================================

with st.sidebar:

    st.markdown(
        """
        <h1 style="margin-bottom:0;">CvScout</h1>
        """,
        unsafe_allow_html=True
    )

    st.markdown(
        "### Data Science Resume Screening"
    )

    st.markdown("---")

    st.markdown(
        """
        **Data Science Pipeline**

        • Resume Parsing

        • NLP Preprocessing

        • TF-IDF

        • Cosine Similarity

        • Sentence Transformers

        • Feature Engineering

        • Candidate Scoring
        """
    )

    st.markdown("---")

    st.caption(
        "Free/open-source machine learning and NLP pipeline"
    )


# ============================================================
# Main header
# ============================================================

st.title("CvScout")

st.caption(
    "Data Science Resume Screening System"
)


st.markdown(
    """
    Upload one Job Description and multiple CVs. CvScout extracts candidate information,
    compares job-related text and skills, and ranks the CVs using NLP and similarity features.
    Scores support human review; they are not probabilities or validated hiring decisions.
    """
)


# ============================================================
# Upload section
# ============================================================

col1, col2 = st.columns(2)

with col1:

    jd_file = st.file_uploader(
        "Upload Job Description",
        type=["pdf", "docx"],
        key="jd_upload"
    )


with col2:

    resume_files = st.file_uploader(
        "Upload Resumes",
        type=["pdf", "docx"],
        accept_multiple_files=True,
        key="resume_upload"
    )


# ============================================================
# Job description
# ============================================================

jd_text = ""

if jd_file:

    jd_text = extract_text(
        jd_file
    )

    if jd_text.strip() and not jd_text.startswith("ERROR_READING_"):

        st.success(
            "Job Description loaded successfully."
        )

        with st.expander(
            "Preview Job Description"
        ):

            st.text_area(
                "JD Text",
                jd_text[:10000],
                height=250,
                label_visibility="collapsed"
            )

    else:

        st.error(
            "Could not extract text from the Job Description."
        )


# ============================================================
# Resume upload confirmation
# ============================================================

if resume_files:

    st.success(
        f"{len(resume_files)} resume(s) uploaded."
    )


# ============================================================
# Analyze
# ============================================================

analyze = st.button(
    "🚀 Analyze Resumes",
    use_container_width=True,
    type="primary"
)


# ============================================================
# Analysis
# ============================================================

if analyze:

    if not jd_text.strip():

        st.error(
            "Please upload a valid Job Description first."
        )

        st.stop()

    if not resume_files:

        st.error(
            "Please upload at least one resume."
        )

        st.stop()

    # --------------------------------------------------------
    # Load model
    # --------------------------------------------------------

    with st.spinner(
        "Loading Sentence Transformer model... (first run may download the model)"
    ):
        try:
            model = load_embedding_model()
        except Exception as exc:
            st.error(
                "Could not load the Sentence Transformer model. Check the internet connection "
                f"and package installation. Details: {exc}"
            )
            st.stop()

    # --------------------------------------------------------
    # Process
    # --------------------------------------------------------

    start_time = time.time()

    results = []

    progress = st.progress(
        0
    )

    status = st.empty()

    total = len(resume_files)

    for index, resume in enumerate(
        resume_files,
        start=1
    ):

        status.write(
            f"Processing {resume.name} "
            f"({index}/{total})..."
        )

        result = process_candidate(
            resume,
            jd_text,
            model
        )

        results.append(
            result
        )

        progress.progress(
            index / total
        )

    elapsed = round(
        time.time() - start_time,
        1
    )

    progress.empty()
    status.empty()

    # --------------------------------------------------------
    # DataFrame
    # --------------------------------------------------------

    df = pd.DataFrame(
        results
    )

    df = df[
        EXPORT_COLUMNS
    ]

    df = df.sort_values(
        by="Relevancy (%)",
        ascending=False
    ).reset_index(
        drop=True
    )

    # --------------------------------------------------------
    # Save in session
    # --------------------------------------------------------

    st.session_state["results"] = df
    st.session_state["analysis_time"] = elapsed


# ============================================================
# Results
# ============================================================

if "results" in st.session_state:

    df = st.session_state["results"]

    elapsed = st.session_state.get(
        "analysis_time",
        0
    )

    st.markdown("---")

    st.header(
        "Screening Results"
    )

    # --------------------------------------------------------
    # Metrics
    # --------------------------------------------------------

    total_candidates = len(df)
    average_score = round(df["Relevancy (%)"].mean(), 1) if total_candidates else 0
    highest_score = round(df["Relevancy (%)"].max(), 1) if total_candidates else 0

    m1, m2, m3 = st.columns(3)

    with m1:
        st.metric("Candidates Analysed", total_candidates)

    with m2:
        st.metric("Average Relevancy", f"{average_score}%")

    with m3:
        st.metric("Highest Relevancy", f"{highest_score}%")

    st.caption(
        f"Processed in {elapsed} seconds. Relevancy is an estimated alignment score, "
        "not a probability of hiring success or a validated accuracy percentage."
    )

    st.info(
        "Candidates are ranked by estimated alignment with the Job Description. "
        "Review the CV and individual matching signals before making recruitment decisions."
    )

    # --------------------------------------------------------
    # Main result table
    # --------------------------------------------------------

    st.markdown(
        "### Ranked Candidates"
    )

    display_columns = [
        "Relevancy (%)",
            "File Name",
        "Full Name",
        "First Name",
        "Email",
        "Phone",
        "City",
        "Current Job Title",
        "Experience (Years)",
        "Current Company",
        "Skills",
        "Matched Skills",
        "Missing Skills",
        "ATS Score",
        "TF-IDF Similarity",
        "Semantic Similarity",
        "Skill Match",
        "Keyword Match",
        "Experience Match",
        "Education Match"
    ]

    st.dataframe(
        df[display_columns],
        use_container_width=True,
        hide_index=True,
        height=550
    )

    # --------------------------------------------------------
    # Candidate details
    # --------------------------------------------------------

    st.markdown(
        "### Candidate Details"
    )

    selected_file = st.selectbox(
        "Select candidate",
        df["File Name"].tolist()
    )

    selected = df[
        df["File Name"] == selected_file
    ].iloc[0]

    c1, c2, c3 = st.columns(3)

    with c1:

        st.metric(
            "Relevancy",
            f"{selected['Relevancy (%)']}%"
        )

    with c2:

        st.metric(
            "Semantic Similarity",
            f"{selected['Semantic Similarity']}%"
        )

    with c3:

        st.metric(
            "Skill Match",
            f"{selected['Skill Match']}%"
        )

    d1, d2 = st.columns(2)

    with d1:

        st.markdown(
            f"""
            **Candidate:** {selected['Full Name']}

            **Email:** {selected['Email']}

            **Phone:** {selected['Phone']}

            **Current Job:** {selected['Current Job Title']}

            **Company:** {selected['Current Company']}

            **Experience:** {selected['Experience (Years)']} years

            **Location:** {selected['Location']}
            """
        )

    with d2:

        st.markdown(
            f"""
            **Matched Skills:** {selected['Matched Skills']}

            **Missing Skills:** {selected['Missing Skills']}

            **TF-IDF:** {selected['TF-IDF Similarity']}%

            **Keyword Match:** {selected['Keyword Match']}%

            **Experience Match:** {selected['Experience Match']}%

            **Education Match:** {selected['Education Match']}%
            """
        )

    # --------------------------------------------------------
    # Export
    # --------------------------------------------------------

    st.markdown("---")

    st.markdown(
        "### Export Results"
    )

    csv_data = df.to_csv(
        index=False
    ).encode("utf-8")

    excel_data = create_excel(
        df
    )

    e1, e2 = st.columns(2)

    with e1:

        st.download_button(
            "Download CSV",
            data=csv_data,
            file_name="cvscout_results.csv",
            mime="text/csv",
            use_container_width=True
        )

    with e2:

        st.download_button(
            "Download Excel",
            data=excel_data,
            file_name="cvscout_results.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )